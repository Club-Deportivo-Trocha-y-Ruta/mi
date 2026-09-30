"""Pruebas de la plantilla FO-GDD-057 v006 versionada (feature 047, T009).

La plantilla se genera una sola vez con ``scripts/build_imderty_template.py``
a partir del libro del dueño del club, que nunca entra al repo. Estas pruebas
garantizan que el binario versionado no arrastra datos de participantes
(menores de edad, Ley 1581) ni valores del encabezado, y que su estructura es
la que el generador de la planilla (US1) espera.

Nunca se imprime el valor de una celda en un mensaje de aserción: solo la
hoja y la coordenada.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.imderty.barrios_seed import BARRIOS_YUMBO

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "templates"
    / "documents"
    / "imderty"
    / "fo_gdd_057_v006.xlsx"
)

MONTH_SHEETS = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
]

FIRST_DATA_ROW = 24
LAST_DATA_ROW = 504
HEADER_ROW = 23
LAST_COLUMN = 54  # BB

HEADER_VALUE_CELLS = ("D7", "D8", "D9", "D10", "I7")

# Celdas donde se marca la "X" de cada programa (bloque B12:I18).
PROGRAM_MARK_CELLS = ("C14", "D14", "E14", "F14", "H14", "I14", "C18", "E18", "H18")


@pytest.fixture(scope="module")
def workbook():
    assert TEMPLATE_PATH.is_file(), "falta la plantilla versionada"
    return load_workbook(TEMPLATE_PATH)


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def test_sheets_in_calendar_order_plus_sector(workbook):
    assert workbook.sheetnames == [*MONTH_SHEETS, "SECTOR"]


def test_participant_rows_hold_only_formulas_or_blanks(workbook):
    offenders: list[str] = []
    for name in MONTH_SHEETS:
        ws = workbook[name]
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=ws.max_row):
            for cell in row:
                if cell.value is None or _is_formula(cell.value):
                    continue
                offenders.append(f"{name}!{cell.coordinate}")
    assert not offenders, f"celdas con valor en filas >= 24: {offenders[:10]}"


def test_participant_rows_keep_their_formulas(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        for row in (FIRST_DATA_ROW, 250, LAST_DATA_ROW):
            for col in ("F", "R", "AZ", "BA", "BB"):
                assert _is_formula(ws[f"{col}{row}"].value), f"{name}!{col}{row}"
        assert "DATEDIF" in ws["F100"].value
        assert "VLOOKUP" in ws["R100"].value
        for col in ("AZ", "BA", "BB"):
            assert "COUNTIF" in ws[f"{col}100"].value


def test_header_values_are_blank(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        for ref in HEADER_VALUE_CELLS:
            assert ws[ref].value is None, f"{name}!{ref} no está vacía"


def test_there_are_no_program_x_marks(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        for ref in PROGRAM_MARK_CELLS:
            assert ws[ref].value is None, f"{name}!{ref} no está vacía"
        for row in ws.iter_rows(min_row=1, max_row=HEADER_ROW):
            for cell in row:
                if isinstance(cell.value, str):
                    assert cell.value.strip().upper() != "X", (
                        f"marca X en {name}!{cell.coordinate}"
                    )


def test_program_labels_are_kept(workbook):
    ws = workbook["ENERO"]
    labels = {
        "C12": "MASIFICACIÓN",
        "D13": "PRIMERA INFANCIA",
        "E13": "HEVS",
        "F13": "RECREACIÓN",
        "I13": "LECYD-CDA",
        "C16": "COMPETENCIA",
        "C17": "CONJUNTO",
        "E17": "INDIVIDUAL",
        "H17": "ADAPTADO",
    }
    for ref, expected in labels.items():
        assert ws[ref].value.strip() == expected, ref


def test_table_names_are_unique(workbook):
    names: list[str] = []
    for name in MONTH_SHEETS:
        tables = workbook[name].tables
        assert len(tables) == 1, name
        (table,) = tables.values()
        assert table.ref == f"A{HEADER_ROW}:BB{LAST_DATA_ROW}"
        names.append(table.displayName)
    assert len(names) == len(set(names))
    assert names == [f"Tabla{m.capitalize()}" for m in MONTH_SHEETS]


def test_table_ids_are_unique(workbook):
    ids = [next(iter(workbook[n].tables.values())).id for n in MONTH_SHEETS]
    assert len(ids) == len(set(ids))


def test_day_headers_equal_table_column_names(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        (table,) = ws.tables.values()
        header = [
            ws[f"{get_column_letter(i)}{HEADER_ROW}"].value
            for i in range(1, LAST_COLUMN + 1)
        ]
        columns = [col.name for col in table.tableColumns]
        assert header == columns, name


def test_formulas_reference_only_their_own_table(workbook):
    all_names = {next(iter(workbook[n].tables.values())).displayName for n in MONTH_SHEETS}
    for name in MONTH_SHEETS:
        ws = workbook[name]
        own = next(iter(ws.tables.values())).displayName
        for row in ws.iter_rows(min_row=1, max_row=LAST_DATA_ROW):
            for cell in row:
                if not _is_formula(cell.value):
                    continue
                assert "Tabla13" not in cell.value, f"{name}!{cell.coordinate}"
                for other in all_names - {own}:
                    assert f"{other}[" not in cell.value, f"{name}!{cell.coordinate}"


def test_title_block_code_and_version(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        assert ws["C2"].value.strip() == "INSTITUTO MUNICIPAL DE DEPORTE Y RECREACIÓN DE YUMBO"
        assert "GESTIÓN DEL DEPORTE" in ws["C3"].value
        assert ws["H3"].value.strip() == "CÓDIGO"
        assert ws["I3"].value.strip() == "FO-GDD-057"
        assert ws["H4"].value.strip() == "VERSIÓN"
        assert str(ws["I4"].value).strip() == "006"


def test_barrio_dropdown_is_a_standard_validation(workbook):
    for name in MONTH_SHEETS:
        ws = workbook[name]
        matches = [
            dv
            for dv in ws.data_validations.dataValidation
            if "Q24:Q504" in str(dv.sqref)
        ]
        assert len(matches) == 1, name
        assert matches[0].type == "list"
        assert matches[0].formula1 == f"SECTOR!$A$2:$A${len(BARRIOS_YUMBO) + 1}"


def test_sector_sheet_matches_seed(workbook):
    ws = workbook["SECTOR"]
    rows = [
        (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
        for r in range(2, ws.max_row + 1)
    ]
    expected = [
        (barrio, int(zone) if zone.isdigit() else zone) for barrio, zone in BARRIOS_YUMBO
    ]
    assert rows == expected


def test_package_has_no_calc_chain_printer_settings_or_defined_names():
    with zipfile.ZipFile(TEMPLATE_PATH) as zf:
        names = zf.namelist()
        assert not any("calcChain" in n for n in names)
        assert not any("printerSettings" in n for n in names)
        workbook_xml = zf.read("xl/workbook.xml").decode("utf-8")
        assert "#REF!" not in workbook_xml
        assert "absPath" not in workbook_xml
        assert 'fullCalcOnLoad="1"' in workbook_xml
        media = [n for n in names if n.startswith("xl/media/")]
        assert len(media) == 1


def test_each_month_sheet_keeps_the_logo(workbook):
    for name in MONTH_SHEETS:
        assert len(workbook[name]._images) == 1, name


def test_openpyxl_round_trip_keeps_structure(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    reloaded = load_workbook(buffer)
    assert reloaded.sheetnames == [*MONTH_SHEETS, "SECTOR"]
    for name in MONTH_SHEETS:
        ws = reloaded[name]
        assert len(ws.tables) == 1
        assert len(ws.data_validations.dataValidation) == 12
        assert len(ws._images) == 1
