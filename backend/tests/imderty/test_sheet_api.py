"""T016 (specs/047-imderty-attendance-sheet/tasks.md, US1):
``POST /api/clubs/{club_id}/imderty-sheet`` with ``from == to``.

Covers: a 200 xlsx response with the contract's filename shape; 403 for a
parent, an athlete account and a foreign coach; an ``export`` audit row that
carries only counts; nothing written to disk; and no synthetic athlete name
reaching any log line.

The export audit row carries exactly the contract's keys
(``contracts/api.md``): ``{document_kind, from_month, to_month, row_count,
gap_count}`` — all in ``services/audit.py::META_ALLOWLIST`` since the W3
stories gate. Columns L/N/O carry the sensitive block only for an athlete
with an active authorization; column M is always blank.

The engine adds, on top of ``conftest``'s shared ``_TABLES``, every table
``attendance_grid.py``/``workbook.py`` touch — none of them are in the
shared subset, and this file must not edit ``conftest.py`` (same idiom as
``test_settings_api.py``/``test_imderty_profile_api.py``).

Every name here is fictitious (CLAUDE.md, Ley 1581).
"""
from __future__ import annotations

import logging

import pytest
import pytest_asyncio
from sqlalchemy import select

from app.models import Base
from app.models.audit_log import AuditLog
from app.models.user import UserRole
from app.services.audit import AuditAction, AuditEntityType
from tests.fixtures.race_history_fixtures import create_user
from tests.imderty.conftest import ImdertyScenario

_SHEET_TABLES = (
    "imderty_barrios",
    "athlete_imderty_profiles",
    "club_imderty_settings",
    "training_sessions",
    "session_attendance",
    "calendar_events",
    "event_audiences",
    "event_attendances",
    "race_events",
    "race_results",
    "athlete_sensitive_authorizations",
    "athlete_sensitive_data",
)

_ATHLETE_ONE_FIRST_NAME = "Atleta Ficticio Uno"
_ATHLETE_ONE_LAST_NAME = "Ficticio"
_ATHLETE_ACCOUNT_USER_ID = 4753


@pytest_asyncio.fixture(autouse=True)
async def _ensure_sheet_tables(imderty_engine):
    """Create the extra tables this route needs on the same in-memory
    engine, without touching ``conftest.py`` (owned by a foundation task,
    not this one)."""
    async with imderty_engine.begin() as conn:
        for name in _SHEET_TABLES:
            await conn.run_sync(
                lambda c, name=name: Base.metadata.tables[name].create(
                    c, checkfirst=True
                )
            )
    yield


@pytest_asyncio.fixture
async def seeded(imderty_scenario: ImdertyScenario) -> ImdertyScenario:
    """Scenario plus a login user with ``role=athlete`` for one of the
    seeded athletes, used only for the "athlete account" 403 case."""
    s = imderty_scenario.session
    await create_user(
        s,
        user_id=_ATHLETE_ACCOUNT_USER_ID,
        role=UserRole.athlete,
        first_name=_ATHLETE_ONE_FIRST_NAME,
        last_name=_ATHLETE_ONE_LAST_NAME,
        can_login=False,
    )
    await s.commit()
    return imderty_scenario


def _sheet_url(club_id: int) -> str:
    return f"/api/clubs/{club_id}/imderty-sheet"


_BODY_AUG_2026 = {"from": "2026-08", "to": "2026-08"}


async def _audit_rows(scenario: ImdertyScenario) -> list[AuditLog]:
    result = await scenario.session.execute(
        select(AuditLog)
        .where(AuditLog.entity_type == AuditEntityType.imderty_attendance_sheet.value)
        .order_by(AuditLog.id)
        .execution_options(populate_existing=True)
    )
    return list(result.scalars().all())


# ---------------------------------------------------------------------------
# 200 — xlsx response
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_sheet_returns_xlsx(seeded, admin_client):
    response = await admin_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    disposition = response.headers["content-disposition"]
    assert disposition == (
        'attachment; filename="FO-GDD-057_asistencia_2026-08_2026-08.xlsx"'
    )
    # The filename never carries a name, a date of birth or any athlete id.
    assert _ATHLETE_ONE_FIRST_NAME.lower() not in disposition.lower()
    assert int(response.headers["content-length"]) == len(response.content)
    assert len(response.content) > 0


@pytest.mark.asyncio
async def test_generate_sheet_coach_of_the_club_succeeds(seeded, coach_client):
    response = await coach_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 200


# ---------------------------------------------------------------------------
# 403 — parent, athlete account, foreign coach
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_sheet_parent_forbidden(seeded, parent_client):
    response = await parent_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_generate_sheet_foreign_coach_forbidden(seeded, foreign_coach_client):
    response = await foreign_coach_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_generate_sheet_athlete_account_forbidden(
    seeded, imderty_client_factory
):
    async with imderty_client_factory(_ATHLETE_ACCOUNT_USER_ID) as client:
        response = await client.post(_sheet_url(seeded.club_id), json=_BODY_AUG_2026)
    assert response.status_code == 403


# ---------------------------------------------------------------------------
# Audit — export row with only counts
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_sheet_records_export_audit(seeded, admin_client):
    response = await admin_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 200

    rows = await _audit_rows(seeded)
    assert len(rows) == 1
    entry = rows[0]
    assert entry.action == AuditAction.export.value
    assert entry.club_id == seeded.club_id
    assert entry.athlete_id is None
    assert entry.diff_json is None

    meta = entry.meta_json or {}
    # Exactly the contract's keys (contracts/api.md, data-model.md §audit).
    assert set(meta.keys()) == {
        "document_kind",
        "from_month",
        "to_month",
        "row_count",
        "gap_count",
    }
    assert meta["document_kind"] == "imderty_attendance_xlsx"
    assert meta["from_month"] == "2026-08"
    assert meta["to_month"] == "2026-08"
    # The three scenario athletes are active in August 2026 and none of
    # them has an IMDERTY profile yet, so every one of them has gaps.
    assert meta["row_count"] == 3
    assert meta["gap_count"] == 3

    # No athlete name, surname or any other value ever reaches meta_json.
    dumped = str(meta)
    assert _ATHLETE_ONE_FIRST_NAME.lower() not in dumped.lower()
    assert _ATHLETE_ONE_LAST_NAME.lower() not in dumped.lower()


# ---------------------------------------------------------------------------
# Not stored — nothing is written to disk
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_sheet_does_not_write_to_disk(seeded, admin_client, monkeypatch):
    from pathlib import Path

    def _forbidden_write(self, *args, **kwargs):  # noqa: ANN001, ANN002, ANN003
        raise AssertionError(
            "The IMDERTY sheet route must not write any file to disk"
        )

    monkeypatch.setattr(Path, "write_bytes", _forbidden_write)
    monkeypatch.setattr(Path, "write_text", _forbidden_write)

    response = await admin_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 200


# ---------------------------------------------------------------------------
# Logs — no synthetic athlete name ever logged
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_generate_sheet_never_logs_athlete_names(
    seeded, admin_client, caplog
):
    with caplog.at_level(logging.DEBUG):
        response = await admin_client.post(
            _sheet_url(seeded.club_id), json=_BODY_AUG_2026
        )
    assert response.status_code == 200

    for record in caplog.records:
        message = record.getMessage().lower()
        assert _ATHLETE_ONE_FIRST_NAME.lower() not in message
        assert _ATHLETE_ONE_LAST_NAME.lower() not in message


# ---------------------------------------------------------------------------
# US4 — period of months (T047)
# ---------------------------------------------------------------------------


def _open_xlsx(content: bytes):  # noqa: ANN202
    from io import BytesIO

    from openpyxl import load_workbook

    return load_workbook(BytesIO(content))


def _first_names(ws) -> list[str]:  # noqa: ANN001
    from app.services.imderty.workbook import FIRST_DATA_ROW, MAX_ROWS

    return [
        str(ws[f"B{r}"].value)
        for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + MAX_ROWS)
        if ws[f"B{r}"].value is not None
    ]


@pytest.mark.asyncio
async def test_period_across_year_boundary_order_and_filename(seeded, admin_client):
    response = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-11", "to": "2027-02"}
    )
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        'attachment; filename="FO-GDD-057_asistencia_2026-11_2027-02.xlsx"'
    )
    wb = _open_xlsx(response.content)
    assert wb.sheetnames == ["NOVIEMBRE", "DICIEMBRE", "ENERO", "FEBRERO", "SECTOR"]

    rows = await _audit_rows(seeded)
    assert len(rows) == 1
    meta = rows[0].meta_json or {}
    assert meta["from_month"] == "2026-11"
    assert meta["to_month"] == "2027-02"
    # Three athletes in each of the four month sheets.
    assert meta["row_count"] == 12


@pytest.mark.asyncio
async def test_period_sheets_match_single_month_downloads(seeded, admin_client):
    period = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-08", "to": "2026-10"}
    )
    assert period.status_code == 200
    period_wb = _open_xlsx(period.content)
    assert period_wb.sheetnames == ["AGOSTO", "SEPTIEMBRE", "OCTUBRE", "SECTOR"]

    for month, name in (
        ("2026-08", "AGOSTO"),
        ("2026-09", "SEPTIEMBRE"),
        ("2026-10", "OCTUBRE"),
    ):
        single = await admin_client.post(
            _sheet_url(seeded.club_id), json={"from": month, "to": month}
        )
        assert single.status_code == 200
        single_ws = _open_xlsx(single.content)[name]
        period_ws = period_wb[name]
        assert [tuple(c.value for c in r) for r in period_ws.iter_rows()] == [
            tuple(c.value for c in r) for r in single_ws.iter_rows()
        ], name


@pytest.mark.asyncio
async def test_period_membership_is_per_month(seeded, admin_client):
    from datetime import date

    from app.models.athlete import Athlete

    athlete = await seeded.session.get(Athlete, seeded.parent_athlete_one_id)
    assert athlete is not None
    athlete.club_join_date = date(2026, 9, 10)
    await seeded.session.commit()

    response = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-08", "to": "2026-09"}
    )
    assert response.status_code == 200
    wb = _open_xlsx(response.content)
    joined = _ATHLETE_ONE_FIRST_NAME.upper()
    august = _first_names(wb["AGOSTO"])
    september = _first_names(wb["SEPTIEMBRE"])
    assert joined not in august
    assert joined in september
    assert len(september) == len(august) + 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "body",
    [
        {"from": "2026-01", "to": "2027-01"},  # 13 months
        {"from": "2026-09", "to": "2026-08"},  # to < from
        {"from": "2026-13", "to": "2027-01"},  # bad month
    ],
)
async def test_period_invalid_range_rejected(seeded, admin_client, body):
    response = await admin_client.post(_sheet_url(seeded.club_id), json=body)
    assert response.status_code == 422
    assert await _audit_rows(seeded) == []


@pytest.mark.asyncio
async def test_period_twelve_months_accepted(seeded, admin_client):
    response = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-02", "to": "2027-01"}
    )
    assert response.status_code == 200
    wb = _open_xlsx(response.content)
    assert len(wb.sheetnames) == 13
    assert wb.sheetnames[0] == "FEBRERO"
    assert wb.sheetnames[-2:] == ["ENERO", "SECTOR"]


@pytest.mark.asyncio
async def test_readiness_period_range(seeded, admin_client):
    url = f"/api/clubs/{seeded.club_id}/imderty-sheet/readiness"
    ok = await admin_client.get(url, params={"from": "2026-11", "to": "2027-02"})
    assert ok.status_code == 200
    assert ok.json()["months"] == ["2026-11", "2026-12", "2027-01", "2027-02"]

    too_long = await admin_client.get(url, params={"from": "2026-01", "to": "2027-01"})
    assert too_long.status_code == 422
    assert "2027-01" not in str(too_long.json())


@pytest.mark.asyncio
async def test_period_filename_is_canonical(seeded, admin_client):
    """Only strict ``YYYY-MM`` is accepted, so the file name and the audit
    meta can only ever hold canonical months."""
    response = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-8", "to": "2026-09"}
    )
    assert response.status_code == 422
    assert await _audit_rows(seeded) == []

    ok = await admin_client.post(
        _sheet_url(seeded.club_id), json={"from": "2026-08", "to": "2026-09"}
    )
    assert ok.headers["content-disposition"] == (
        'attachment; filename="FO-GDD-057_asistencia_2026-08_2026-09.xlsx"'
    )


# ---------------------------------------------------------------------------
# US2 — sensitive columns L/N/O only under an active authorization (T032/T038)
# ---------------------------------------------------------------------------


def _row_of(ws, first_name: str) -> int:  # noqa: ANN001
    from app.services.imderty.workbook import FIRST_DATA_ROW, MAX_ROWS

    for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + MAX_ROWS):
        if ws[f"B{r}"].value == first_name.upper():
            return r
    raise AssertionError("athlete row not found in the sheet")


@pytest.mark.asyncio
async def test_sheet_fills_sensitive_columns_only_when_authorized(
    seeded, coach_client
):
    authorized = seeded.parent_athlete_one_id
    created = await coach_client.post(
        f"/api/athletes/{authorized}/sensitive-authorizations",
        json={"guardian_user_id": seeded.parent_user_id, "authorized_on": "2026-08-01"},
    )
    assert created.status_code == 201
    saved = await coach_client.put(
        f"/api/athletes/{authorized}/sensitive-data",
        json={"ethnicity": "RAIZAL", "disability": "VISUAL", "conflict_victim": "si"},
    )
    assert saved.status_code == 200

    response = await coach_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 200
    ws = _open_xlsx(response.content)["AGOSTO"]

    row = _row_of(ws, _ATHLETE_ONE_FIRST_NAME)
    assert ws[f"L{row}"].value == "VISUAL"
    assert ws[f"N{row}"].value == "SI"
    assert ws[f"O{row}"].value == "RAIZAL"
    assert ws[f"M{row}"].value is None

    for other in ("Atleta Ficticio Dos", "Atleta Ficticio Tres"):
        other_row = _row_of(ws, other)
        for column in ("L", "M", "N", "O"):
            assert ws[f"{column}{other_row}"].value is None, (other, column)


@pytest.mark.asyncio
async def test_sheet_blanks_sensitive_columns_after_withdrawal(seeded, coach_client):
    athlete_id = seeded.parent_athlete_one_id
    created = await coach_client.post(
        f"/api/athletes/{athlete_id}/sensitive-authorizations",
        json={"guardian_user_id": seeded.parent_user_id, "authorized_on": "2026-08-01"},
    )
    assert created.status_code == 201
    saved = await coach_client.put(
        f"/api/athletes/{athlete_id}/sensitive-data",
        json={"ethnicity": "RAIZAL", "disability": "VISUAL", "conflict_victim": "si"},
    )
    assert saved.status_code == 200
    withdrawn = await coach_client.post(
        f"/api/athletes/{athlete_id}/sensitive-authorizations/withdraw"
    )
    assert withdrawn.status_code == 200

    response = await coach_client.post(
        _sheet_url(seeded.club_id), json=_BODY_AUG_2026
    )
    assert response.status_code == 200
    ws = _open_xlsx(response.content)["AGOSTO"]
    row = _row_of(ws, _ATHLETE_ONE_FIRST_NAME)
    for column in ("L", "M", "N", "O"):
        assert ws[f"{column}{row}"].value is None, column
