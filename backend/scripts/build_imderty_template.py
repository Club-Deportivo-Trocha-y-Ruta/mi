"""Construye la plantilla sanitizada FO-GDD-057 v006 (planilla IMDERTY).

Feature 047 (T008, research R2). Script de desarrollo, se ejecuta a mano una
sola vez por revisión del formato oficial:

  cd backend && .venv/bin/python scripts/build_imderty_template.py <libro_del_dueño.xlsx>

El libro del dueño del club (con datos reales de participantes, menores de
edad) se pasa como argumento y NUNCA se copia al repo. El resultado se guarda
en ``templates/documents/imderty/fo_gdd_057_v006.xlsx`` y es lo único que se
versiona.

Qué hace, trabajando directamente sobre las partes XML del paquete (openpyxl
pierde la validación x14 del barrio y no clona imágenes ni tablas):

1. Toma la hoja ``AGOSTO`` como base.
2. Borra todo valor de las filas 24–504 (todas las columnas) y conserva las
   fórmulas: ``F`` = DATEDIF (edad), ``R`` = VLOOKUP (comuna),
   ``AZ:BB`` = COUNTIF (totales A/F/E). Las fórmulas compartidas se expanden
   celda por celda y se descartan todos los valores en caché (derivan de
   datos de participantes).
3. Borra los valores del encabezado (``D7``–``D10``, ``I7``) y toda marca
   "X" de programa; conserva rótulos, estilos, combinaciones, validaciones y
   el logo.
4. Clona la base en 12 hojas ``ENERO``…``DICIEMBRE``, cada una con su tabla
   (``TablaEnero``…, id único), sus referencias estructuradas reescritas, y
   su propio dibujo y relaciones de imagen.
5. Reconstruye ``SECTOR`` desde ``BARRIOS_YUMBO`` y convierte la lista
   desplegable de barrio (x14) en una validación estándar
   ``SECTOR!$A$2:$A$<n>`` que openpyxl sí conserva en tiempo de ejecución.
6. Reconstruye ``sharedStrings`` solo con los textos de la plantilla, sin
   ``calcChain``, sin ``printerSettings``, sin nombres definidos rotos y con
   ``fullCalcOnLoad="1"``.

Privacidad: el script nunca imprime valores de celdas del libro fuente. Si al
verificar la salida encuentra un valor que no sea fórmula en filas >= 24,
termina con código 2 y no escribe la plantilla.
"""

from __future__ import annotations

import argparse
import copy
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.imderty.barrios_seed import BARRIOS_YUMBO  # noqa: E402

OUTPUT_PATH = BACKEND_DIR / "templates" / "documents" / "imderty" / "fo_gdd_057_v006.xlsx"

BASE_SHEET = "AGOSTO"
SECTOR_SHEET = "SECTOR"
MONTHS = (
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
)

FIRST_DATA_ROW = 24
LAST_DATA_ROW = 504
HEADER_VALUE_CELLS = frozenset({"D7", "D8", "D9", "D10", "I7"})
# Bloque de programas (B12:I19): aquí vive cada marca "X".
PROGRAM_BLOCK_ROWS = range(12, 20)
PROGRAM_BLOCK_COLS = range(column_index_from_string("B"), column_index_from_string("I") + 1)
MONTH_TITLE_CELL = "U21"
FIRST_DAY_COL = "U"
LAST_DAY_COL = "AY"
FORMULA_COLUMNS = ("F", "R", "AZ", "BA", "BB")
BARRIO_RANGE = f"Q{FIRST_DATA_ROW}:Q{LAST_DATA_ROW}"
TABLE_PLACEHOLDER = "TablaPLANTILLA"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XR_UID = "{http://schemas.microsoft.com/office/spreadsheetml/2014/revision}uid"
XR3_UID = "{http://schemas.microsoft.com/office/spreadsheetml/2016/revision3}uid"
X14_DV_URI = "{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}"
NS = {"m": MAIN_NS, "r": REL_NS, "p": PKG_REL_NS}

REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
CT_TABLE = "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"
CT_DRAWING = "application/vnd.openxmlformats-officedocument.drawing+xml"

_CELL_REF = re.compile(r"^([A-Z]+)(\d+)$")
# COUNTIF(Tabla13[[#This Row],[SA1]:[LU31]], "F") → COUNTIF(U24:AY24,"F").
# Los nombres de los días se reescriben por mes en tiempo de ejecución, así
# que una referencia estructurada a ellos se rompería; el rango posicional no.
_DAY_RANGE_COUNTIF = re.compile(
    r'COUNTIF\(\s*(?P<table>[A-Za-z_][\w.]*)\[\[#This Row\],\[[^\]]+\]:\[[^\]]+\]\]\s*,\s*"(?P<mark>[AFE])"\s*\)'
)


class TemplateBuildError(RuntimeError):
    """Error estructural. El mensaje nunca incluye valores de celdas."""


def _q(tag: str) -> str:
    return f"{{{MAIN_NS}}}{tag}"


def _split_ref(ref: str) -> tuple[str, int]:
    match = _CELL_REF.match(ref)
    if not match:
        raise TemplateBuildError(f"referencia de celda inesperada: {ref}")
    return match.group(1), int(match.group(2))


def _xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _strip_revision_uids(root: etree._Element) -> None:
    for elem in root.iter():
        for attr in (XR_UID, XR3_UID):
            if attr in elem.attrib:
                del elem.attrib[attr]


# --------------------------------------------------------------------------
# Lectura del paquete fuente
# --------------------------------------------------------------------------


class SourcePackage:
    def __init__(self, path: Path) -> None:
        self._zip = zipfile.ZipFile(path)
        self.names = set(self._zip.namelist())

    def read(self, part: str) -> bytes:
        return self._zip.read(part)

    def xml(self, part: str) -> etree._Element:
        return etree.fromstring(self.read(part))

    def rels(self, part: str) -> dict[str, tuple[str, str]]:
        """rId → (tipo, parte absoluta) para las relaciones de ``part``."""
        folder, name = posixpath.split(part)
        rels_part = posixpath.join(folder, "_rels", f"{name}.rels")
        if rels_part not in self.names:
            return {}
        out: dict[str, tuple[str, str]] = {}
        for rel in self.xml(rels_part).findall("p:Relationship", NS):
            target = posixpath.normpath(posixpath.join(folder, rel.get("Target")))
            out[rel.get("Id")] = (rel.get("Type").rsplit("/", 1)[-1], target)
        return out

    def sheet_parts(self) -> dict[str, str]:
        workbook = self.xml("xl/workbook.xml")
        rels = self.rels("xl/workbook.xml")
        return {
            sheet.get("name"): rels[sheet.get(f"{{{REL_NS}}}id")][1]
            for sheet in workbook.find("m:sheets", NS)
        }


def _related(rels: dict[str, tuple[str, str]], kind: str) -> str:
    targets = [target for rel_kind, target in rels.values() if rel_kind == kind]
    if len(targets) != 1:
        raise TemplateBuildError(f"se esperaba exactamente una relación '{kind}', hay {len(targets)}")
    return targets[0]


# --------------------------------------------------------------------------
# Tabla de textos compartidos nueva (solo textos de la plantilla)
# --------------------------------------------------------------------------


class SharedStrings:
    def __init__(self) -> None:
        self._items: list[etree._Element] = []
        self._index: dict[bytes, int] = {}
        self.references = 0

    def add_item(self, si: etree._Element) -> int:
        item = copy.deepcopy(si)
        key = etree.tostring(item)
        if key not in self._index:
            self._index[key] = len(self._items)
            self._items.append(item)
        self.references += 1
        return self._index[key]

    def add_text(self, text: str) -> int:
        si = etree.Element(_q("si"))
        t = etree.SubElement(si, _q("t"))
        t.text = text
        if text != text.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return self.add_item(si)

    def to_xml(self) -> bytes:
        root = etree.Element(_q("sst"), nsmap={None: MAIN_NS})
        root.set("count", str(self.references))
        root.set("uniqueCount", str(len(self._items)))
        root.extend(self._items)
        return _xml_bytes(root)


def _si_text(si: etree._Element) -> str:
    return "".join(t.text or "" for t in si.iter(_q("t")))


# --------------------------------------------------------------------------
# Hoja mensual base
# --------------------------------------------------------------------------


def _clear_cell(cell: etree._Element) -> None:
    """Deja la celda con su estilo y sin valor."""
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)


def _drop_cached_value(cell: etree._Element) -> None:
    for tag in ("v", "is"):
        for child in cell.findall(_q(tag)):
            cell.remove(child)
    cell.attrib.pop("t", None)


def _rewrite_formula(text: str, *, base_table: str, row: int) -> str:
    def _day_range(match: re.Match[str]) -> str:
        if match.group("table") != base_table:
            raise TemplateBuildError("COUNTIF con una tabla ajena a la hoja base")
        mark = match.group("mark")
        return f'COUNTIF({FIRST_DAY_COL}{row}:{LAST_DAY_COL}{row},"{mark}")'

    text = _DAY_RANGE_COUNTIF.sub(_day_range, text)
    return re.sub(rf"\b{re.escape(base_table)}\[", f"{TABLE_PLACEHOLDER}[", text)


def build_base_month_sheet(
    root: etree._Element,
    *,
    source_strings: list[etree._Element],
    sst: SharedStrings,
    base_table: str,
    sector_last_row: int,
) -> None:
    """Sanitiza en sitio la hoja base. Nunca lee valores para mostrarlos."""
    sheet_data = root.find("m:sheetData", NS)

    # Fórmulas compartidas: maestro por índice ``si``.
    masters: dict[str, tuple[str, str]] = {}
    for cell in sheet_data.iter(_q("c")):
        formula = cell.find(_q("f"))
        if formula is not None and formula.get("t") == "shared" and formula.text:
            masters[formula.get("si")] = (formula.text, cell.get("r"))

    for row in sheet_data:
        row_number = int(row.get("r"))
        for cell in row:
            ref = cell.get("r")
            col, _ = _split_ref(ref)
            formula = cell.find(_q("f"))

            if formula is not None:
                if formula.get("t") == "shared":
                    master_text, origin = masters[formula.get("si")]
                    translated = Translator("=" + master_text, origin=origin).translate_formula(ref)
                    formula.text = translated[1:]
                    for attr in ("t", "ref", "si"):
                        formula.attrib.pop(attr, None)
                formula.text = _rewrite_formula(formula.text, base_table=base_table, row=row_number)
                _drop_cached_value(cell)
                continue

            if row_number >= FIRST_DATA_ROW or ref in HEADER_VALUE_CELLS:
                _clear_cell(cell)
                continue

            value = cell.find(_q("v"))
            if cell.get("t") == "s" and value is not None:
                si = source_strings[int(value.text)]
                in_program_block = (
                    row_number in PROGRAM_BLOCK_ROWS
                    and column_index_from_string(col) in PROGRAM_BLOCK_COLS
                )
                if in_program_block and _si_text(si).strip().upper() == "X":
                    _clear_cell(cell)
                    continue
                if ref == MONTH_TITLE_CELL:
                    continue  # se escribe por mes al clonar
                value.text = str(sst.add_item(si))
            elif cell.get("t") in {"str", "inlineStr", "e"}:
                raise TemplateBuildError(f"tipo de celda no soportado en el encabezado: {ref}")

    leftovers = {
        match
        for formula in sheet_data.iter(_q("f"))
        for match in re.findall(r"\b(Tabla\w*)\[", formula.text or "")
        if match != TABLE_PLACEHOLDER
    }
    if leftovers:
        raise TemplateBuildError("quedaron referencias a tablas no reescritas")

    # Vista: sin celda activa heredada.
    for view in root.iterfind("m:sheetViews/m:sheetView", NS):
        view.attrib.pop("tabSelected", None)
        view.attrib.pop("topLeftCell", None)
        for selection in view.findall(_q("selection")):
            view.remove(selection)
        selection = etree.SubElement(view, _q("selection"))
        selection.set("activeCell", f"A{FIRST_DATA_ROW}")
        selection.set("sqref", f"A{FIRST_DATA_ROW}")

    # Lista desplegable de barrio: de x14 a validación estándar.
    ext_list = root.find("m:extLst", NS)
    if ext_list is not None:
        for ext in ext_list.findall(_q("ext")):
            if ext.get("uri") == X14_DV_URI:
                ext_list.remove(ext)
        if len(ext_list) == 0:
            root.remove(ext_list)
    validations = root.find("m:dataValidations", NS)
    if validations is None:
        raise TemplateBuildError("la hoja base no tiene validaciones de datos")
    for dv in validations.findall(_q("dataValidation")):
        if dv.get("sqref") == BARRIO_RANGE:
            validations.remove(dv)
    barrio = etree.SubElement(validations, _q("dataValidation"))
    for key, val in (
        ("type", "list"),
        ("allowBlank", "1"),
        ("showInputMessage", "1"),
        ("showErrorMessage", "1"),
        ("sqref", BARRIO_RANGE),
    ):
        barrio.set(key, val)
    etree.SubElement(barrio, _q("formula1")).text = f"{SECTOR_SHEET}!$A$2:$A${sector_last_row}"
    validations.set("count", str(len(validations)))

    # Sin printerSettings.
    for page_setup in root.findall(_q("pageSetup")):
        page_setup.attrib.pop(f"{{{REL_NS}}}id", None)

    # Relaciones renumeradas: rId1 = dibujo, rId2 = tabla.
    drawing = root.find("m:drawing", NS)
    table_parts = root.find("m:tableParts", NS)
    if drawing is None or table_parts is None or len(table_parts) != 1:
        raise TemplateBuildError("la hoja base debe tener un dibujo y una tabla")
    drawing.set(f"{{{REL_NS}}}id", "rId1")
    table_parts[0].set(f"{{{REL_NS}}}id", "rId2")

    _strip_revision_uids(root)


def build_month_sheet(base: etree._Element, *, table_name: str, title_index: int, selected: bool) -> bytes:
    root = copy.deepcopy(base)
    for formula in root.iter(_q("f")):
        if formula.text and TABLE_PLACEHOLDER in formula.text:
            formula.text = formula.text.replace(TABLE_PLACEHOLDER, table_name)
    title = root.find(f"m:sheetData/m:row/m:c[@r='{MONTH_TITLE_CELL}']", NS)
    if title is None:
        raise TemplateBuildError(f"falta la celda de título del mes {MONTH_TITLE_CELL}")
    _clear_cell(title)
    title.set("t", "s")
    etree.SubElement(title, _q("v")).text = str(title_index)
    if selected:
        root.find("m:sheetViews/m:sheetView", NS).set("tabSelected", "1")
    return _xml_bytes(root)


def build_month_table(base: etree._Element, *, table_id: int, table_name: str, base_table: str) -> bytes:
    root = copy.deepcopy(base)
    _strip_revision_uids(root)
    root.set("id", str(table_id))
    root.set("name", table_name)
    root.set("displayName", table_name)
    for formula in root.iter(_q("calculatedColumnFormula")):
        text = _rewrite_formula(formula.text, base_table=base_table, row=FIRST_DATA_ROW)
        formula.text = text.replace(TABLE_PLACEHOLDER, table_name)
    return _xml_bytes(root)


# --------------------------------------------------------------------------
# Hoja SECTOR
# --------------------------------------------------------------------------


def build_sector_sheet(
    root: etree._Element, *, source_strings: list[etree._Element], sst: SharedStrings
) -> bytes:
    sheet_data = root.find("m:sheetData", NS)
    source_rows = list(sheet_data)
    if len(source_rows) < 3:
        raise TemplateBuildError("la hoja SECTOR no tiene la estructura esperada")

    def _row_attrs(row: etree._Element) -> dict[str, str]:
        return {k: v for k, v in row.attrib.items() if k not in {"r", "spans"}}

    header_row, data_row = source_rows[0], source_rows[2]
    header_styles = [cell.get("s") for cell in header_row]
    data_styles = [cell.get("s") for cell in data_row]
    header_items = []
    for cell in header_row:
        value = cell.find(_q("v"))
        if cell.get("t") != "s" or value is None:
            raise TemplateBuildError("el encabezado de SECTOR no es texto")
        header_items.append(source_strings[int(value.text)])

    new_data = etree.Element(_q("sheetData"))
    row = etree.SubElement(new_data, _q("row"), r="1", spans="1:2", **_row_attrs(header_row))
    for col, style, si in zip(("A", "B"), header_styles, header_items, strict=True):
        cell = etree.SubElement(row, _q("c"), r=f"{col}1", t="s")
        if style:
            cell.set("s", style)
        etree.SubElement(cell, _q("v")).text = str(sst.add_item(si))

    for offset, (barrio, zone) in enumerate(BARRIOS_YUMBO, start=2):
        row = etree.SubElement(new_data, _q("row"), r=str(offset), spans="1:2", **_row_attrs(data_row))
        name_cell = etree.SubElement(row, _q("c"), r=f"A{offset}", t="s")
        zone_cell = etree.SubElement(row, _q("c"), r=f"B{offset}")
        for cell, style in zip((name_cell, zone_cell), data_styles, strict=True):
            if style:
                cell.set("s", style)
        etree.SubElement(name_cell, _q("v")).text = str(sst.add_text(barrio))
        if zone.isdigit():
            etree.SubElement(zone_cell, _q("v")).text = zone
        else:
            zone_cell.set("t", "s")
            etree.SubElement(zone_cell, _q("v")).text = str(sst.add_text(zone))

    root.replace(sheet_data, new_data)
    last_row = len(BARRIOS_YUMBO) + 1
    root.find("m:dimension", NS).set("ref", f"A1:B{last_row}")
    for sort_state in root.findall(_q("sortState")):
        root.remove(sort_state)
    for view in root.iterfind("m:sheetViews/m:sheetView", NS):
        view.attrib.pop("tabSelected", None)
        view.attrib.pop("topLeftCell", None)
        for selection in view.findall(_q("selection")):
            view.remove(selection)
    for page_setup in root.findall(_q("pageSetup")):
        page_setup.attrib.pop(f"{{{REL_NS}}}id", None)
    _strip_revision_uids(root)
    return _xml_bytes(root)


# --------------------------------------------------------------------------
# Partes del paquete
# --------------------------------------------------------------------------


def _relationships(items: list[tuple[str, str, str]]) -> bytes:
    root = etree.Element(f"{{{PKG_REL_NS}}}Relationships", nsmap={None: PKG_REL_NS})
    for rel_id, rel_type, target in items:
        etree.SubElement(root, f"{{{PKG_REL_NS}}}Relationship", Id=rel_id, Type=rel_type, Target=target)
    return _xml_bytes(root)


def _workbook_xml(sheet_names: list[str]) -> bytes:
    root = etree.Element(_q("workbook"), nsmap={None: MAIN_NS, "r": REL_NS})
    etree.SubElement(root, _q("fileVersion"), appName="xl", lastEdited="7", lowestEdited="7", rupBuild="30513")
    etree.SubElement(root, _q("workbookPr"))
    views = etree.SubElement(root, _q("bookViews"))
    etree.SubElement(
        views, _q("workbookView"), xWindow="-120", yWindow="-120", windowWidth="24240", windowHeight="13020", activeTab="0"
    )
    sheets = etree.SubElement(root, _q("sheets"))
    for index, name in enumerate(sheet_names, start=1):
        sheet = etree.SubElement(sheets, _q("sheet"), name=name, sheetId=str(index))
        sheet.set(f"{{{REL_NS}}}id", f"rId{index}")
    etree.SubElement(root, _q("calcPr"), calcId="191028", fullCalcOnLoad="1")
    return _xml_bytes(root)


def _content_types(image_ext: str, image_ct: str, sheet_count: int, month_count: int) -> bytes:
    ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    root = etree.Element(f"{{{ns}}}Types", nsmap={None: ns})
    for ext, ct in (
        (image_ext, image_ct),
        ("rels", "application/vnd.openxmlformats-package.relationships+xml"),
        ("xml", "application/xml"),
    ):
        etree.SubElement(root, f"{{{ns}}}Default", Extension=ext, ContentType=ct)
    overrides = [
        ("/xl/workbook.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"),
        *[(f"/xl/worksheets/sheet{i}.xml", CT_SHEET) for i in range(1, sheet_count + 1)],
        *[(f"/xl/tables/table{i}.xml", CT_TABLE) for i in range(1, month_count + 1)],
        *[(f"/xl/drawings/drawing{i}.xml", CT_DRAWING) for i in range(1, month_count + 1)],
        ("/xl/theme/theme1.xml", "application/vnd.openxmlformats-officedocument.theme+xml"),
        ("/xl/styles.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"),
        ("/xl/sharedStrings.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"),
        ("/docProps/core.xml", "application/vnd.openxmlformats-package.core-properties+xml"),
        ("/docProps/app.xml", "application/vnd.openxmlformats-officedocument.extended-properties+xml"),
    ]
    for part, ct in overrides:
        etree.SubElement(root, f"{{{ns}}}Override", PartName=part, ContentType=ct)
    return _xml_bytes(root)


CORE_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" \
xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" \
xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\
<dc:title>FO-GDD-057 v006 - Planilla de asistencia</dc:title><dc:creator></dc:creator>\
<cp:lastModifiedBy></cp:lastModifiedBy>\
<dcterms:created xsi:type="dcterms:W3CDTF">2015-06-05T18:19:34Z</dcterms:created>\
<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-28T00:00:00Z</dcterms:modified>\
</cp:coreProperties>"""

APP_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" \
xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">\
<Application>Microsoft Excel</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>\
<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>\
<AppVersion>16.0300</AppVersion></Properties>"""

ROOT_RELS = _relationships(
    [
        ("rId1", f"{REL_TYPE}/officeDocument", "xl/workbook.xml"),
        ("rId2", "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties", "docProps/core.xml"),
        ("rId3", f"{REL_TYPE}/extended-properties", "docProps/app.xml"),
    ]
)


def build_template(source_path: Path) -> dict[str, bytes]:
    """Devuelve las partes del paquete de salida (nombre → bytes)."""
    src = SourcePackage(source_path)
    sheet_parts = src.sheet_parts()
    for required in (BASE_SHEET, SECTOR_SHEET):
        if required not in sheet_parts:
            raise TemplateBuildError(f"el libro fuente no tiene la hoja {required}")

    base_part = sheet_parts[BASE_SHEET]
    base_rels = src.rels(base_part)
    table_part = _related(base_rels, "table")
    drawing_part = _related(base_rels, "drawing")
    image_part = _related(src.rels(drawing_part), "image")
    image_ext = posixpath.splitext(image_part)[1].lstrip(".")
    if image_ext != "png":
        raise TemplateBuildError("se esperaba el logo en PNG")

    source_strings = src.xml("xl/sharedStrings.xml").findall("m:si", NS)
    sst = SharedStrings()

    base_table_root = src.xml(table_part)
    base_table = base_table_root.get("displayName")
    sector_last_row = len(BARRIOS_YUMBO) + 1

    base_sheet = src.xml(base_part)
    build_base_month_sheet(
        base_sheet,
        source_strings=source_strings,
        sst=sst,
        base_table=base_table,
        sector_last_row=sector_last_row,
    )

    parts: dict[str, bytes] = {}
    drawing_bytes = src.read(drawing_part)
    image_target = f"../media/{posixpath.basename(image_part)}"
    for index, month in enumerate(MONTHS, start=1):
        table_name = f"Tabla{month.capitalize()}"
        title_index = sst.add_text(month)
        parts[f"xl/worksheets/sheet{index}.xml"] = build_month_sheet(
            base_sheet, table_name=table_name, title_index=title_index, selected=index == 1
        )
        parts[f"xl/worksheets/_rels/sheet{index}.xml.rels"] = _relationships(
            [
                ("rId1", f"{REL_TYPE}/drawing", f"../drawings/drawing{index}.xml"),
                ("rId2", f"{REL_TYPE}/table", f"../tables/table{index}.xml"),
            ]
        )
        parts[f"xl/tables/table{index}.xml"] = build_month_table(
            base_table_root, table_id=index, table_name=table_name, base_table=base_table
        )
        parts[f"xl/drawings/drawing{index}.xml"] = drawing_bytes
        parts[f"xl/drawings/_rels/drawing{index}.xml.rels"] = _relationships(
            [("rId1", f"{REL_TYPE}/image", image_target)]
        )

    sector_index = len(MONTHS) + 1
    parts[f"xl/worksheets/sheet{sector_index}.xml"] = build_sector_sheet(
        src.xml(sheet_parts[SECTOR_SHEET]), source_strings=source_strings, sst=sst
    )

    sheet_names = [*MONTHS, SECTOR_SHEET]
    workbook_rels = [
        (f"rId{i}", f"{REL_TYPE}/worksheet", f"worksheets/sheet{i}.xml") for i in range(1, sector_index + 1)
    ]
    workbook_rels += [
        (f"rId{sector_index + 1}", f"{REL_TYPE}/theme", "theme/theme1.xml"),
        (f"rId{sector_index + 2}", f"{REL_TYPE}/styles", "styles.xml"),
        (f"rId{sector_index + 3}", f"{REL_TYPE}/sharedStrings", "sharedStrings.xml"),
    ]
    parts["xl/workbook.xml"] = _workbook_xml(sheet_names)
    parts["xl/_rels/workbook.xml.rels"] = _relationships(workbook_rels)
    parts["xl/sharedStrings.xml"] = sst.to_xml()
    parts["xl/styles.xml"] = src.read("xl/styles.xml")
    parts["xl/theme/theme1.xml"] = src.read("xl/theme/theme1.xml")
    parts[f"xl/media/{posixpath.basename(image_part)}"] = src.read(image_part)
    parts["docProps/core.xml"] = CORE_XML
    parts["docProps/app.xml"] = APP_XML
    parts["_rels/.rels"] = ROOT_RELS
    parts["[Content_Types].xml"] = _content_types(image_ext, "image/png", sector_index, len(MONTHS))
    return parts


def write_package(parts: dict[str, bytes], path: Path) -> None:
    order = ["[Content_Types].xml", "_rels/.rels", *sorted(k for k in parts if k not in {"[Content_Types].xml", "_rels/.rels"})]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name in order:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            zf.writestr(info, parts[name])


# --------------------------------------------------------------------------
# Verificación (solo coordenadas en los mensajes, nunca valores)
# --------------------------------------------------------------------------


def verify_output(path: Path) -> list[str]:
    problems: list[str] = []

    with zipfile.ZipFile(path) as zf:
        for index, month in enumerate(MONTHS, start=1):
            root = etree.fromstring(zf.read(f"xl/worksheets/sheet{index}.xml"))
            formula_cells: set[str] = set()
            for row in root.find("m:sheetData", NS):
                row_number = int(row.get("r"))
                for cell in row:
                    has_formula = cell.find(_q("f")) is not None
                    has_value = cell.find(_q("v")) is not None or cell.find(_q("is")) is not None
                    if row_number >= FIRST_DATA_ROW:
                        if has_value:
                            problems.append(f"{month}!{cell.get('r')}: valor en fila de participantes")
                        if has_formula:
                            formula_cells.add(cell.get("r"))
                    elif cell.get("r") in HEADER_VALUE_CELLS and has_value:
                        problems.append(f"{month}!{cell.get('r')}: valor de encabezado")
            for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
                for col in FORMULA_COLUMNS:
                    if f"{col}{row_number}" not in formula_cells:
                        problems.append(f"{month}!{col}{row_number}: falta la fórmula")

    workbook = load_workbook(path)
    if workbook.sheetnames != [*MONTHS, SECTOR_SHEET]:
        problems.append("orden de hojas inesperado")
    table_names: list[str] = []
    for month in MONTHS:
        ws = workbook[month]
        for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=ws.max_row):
            for cell in row:
                value = cell.value
                if value is not None and not (isinstance(value, str) and value.startswith("=")):
                    problems.append(f"{month}!{cell.coordinate}: valor no fórmula (openpyxl)")
        for row in ws.iter_rows(min_row=1, max_row=FIRST_DATA_ROW - 1):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().upper() == "X":
                    problems.append(f"{month}!{cell.coordinate}: marca X")
        table_names.extend(table.displayName for table in ws.tables.values())
        if len(ws._images) != 1:
            problems.append(f"{month}: el logo no se conserva")
    if len(table_names) != len(MONTHS) or len(set(table_names)) != len(MONTHS):
        problems.append("nombres de tabla repetidos o faltantes")
    return problems


def verify_with_libreoffice(path: Path) -> str | None:
    """Convierte con LibreOffice headless. Devuelve un error o None."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("LibreOffice no está disponible; se omite esa verificación.")
        return None
    with tempfile.TemporaryDirectory() as out_dir, tempfile.TemporaryDirectory() as profile:
        result = subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--norestore",
                "--convert-to",
                "xlsx",
                "--outdir",
                out_dir,
                str(path),
            ],
            capture_output=True,
            timeout=180,
            check=False,
        )
        converted = Path(out_dir) / path.name
        if result.returncode != 0 or not converted.is_file():
            return f"LibreOffice no pudo abrir la plantilla (código {result.returncode})"
    print("LibreOffice abrió y convirtió la plantilla sin errores.")
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("source", type=Path, help="libro .xlsx del dueño del club (no se versiona)")
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH)
    parser.add_argument("--skip-libreoffice", action="store_true")
    args = parser.parse_args(argv)

    if not args.source.is_file():
        print("No existe el libro fuente indicado.", file=sys.stderr)
        return 1

    try:
        parts = build_template(args.source)
    except TemplateBuildError as exc:
        print(f"Error de estructura: {exc}", file=sys.stderr)
        return 1

    with tempfile.TemporaryDirectory() as tmp:
        candidate = Path(tmp) / args.output.name
        write_package(parts, candidate)

        problems = verify_output(candidate)
        if problems:
            print(f"La plantilla NO es segura ({len(problems)} problemas); no se escribe.", file=sys.stderr)
            for problem in problems[:20]:
                print(f"  - {problem}", file=sys.stderr)
            return 2

        if not args.skip_libreoffice:
            error = verify_with_libreoffice(candidate)
            if error:
                print(error, file=sys.stderr)
                return 3

        args.output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(candidate, args.output)

    print(f"Plantilla escrita: {args.output.relative_to(BACKEND_DIR) if args.output.is_relative_to(BACKEND_DIR) else args.output}")
    print(f"Hojas: {', '.join([*MONTHS, SECTOR_SHEET])}; barrios en SECTOR: {len(BARRIOS_YUMBO)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
